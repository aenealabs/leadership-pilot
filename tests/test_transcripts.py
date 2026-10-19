import io

import pytest
from docx import Document

from leadership_pilot.transcripts import TranscriptError, notes_from_file, vtt_to_notes

TEAMS_VTT = """WEBVTT

0f1c2a3b-1/12-0
00:00:01.200 --> 00:00:04.900
<v Priya Raman>Our ECOs take three weeks</v>

0f1c2a3b-1/13-0
00:00:05.000 --> 00:00:08.100
<v Priya Raman>and nobody can tell me where they are.</v>

0f1c2a3b-2/2-0
00:00:08.400 --> 00:00:10.000
<v Jordan Lee>What would good look like?</v>
"""


def test_teams_vtt_becomes_one_line_per_turn():
    notes = vtt_to_notes(TEAMS_VTT)
    assert notes == ("Priya Raman: Our ECOs take three weeks and nobody can tell me where "
                     "they are.\n\nJordan Lee: What would good look like?")
    assert "-->" not in notes and "00:00" not in notes


def test_plain_vtt_with_speaker_prefixes():
    text = "WEBVTT\n\n1\n00:01.000 --> 00:02.000\nMarcus Bell: Suppliers build to old revs.\n"
    assert notes_from_file("call.vtt", text.encode()) == ("Marcus Bell: Suppliers build to "
                                                          "old revs.")


def test_docx_and_text_files():
    doc = Document()
    doc.add_paragraph("Elena said the BOM reaches SAP a day late.")
    doc.add_paragraph("")
    doc.add_paragraph("Plants find failures on the line.")
    buffer = io.BytesIO()
    doc.save(buffer)
    assert notes_from_file("notes.docx", buffer.getvalue()) == (
        "Elena said the BOM reaches SAP a day late.\n\nPlants find failures on the line.")
    assert notes_from_file("notes.txt", "﻿  hello  \n".encode()) == "hello"


def test_rejects_unreadable_files():
    with pytest.raises(TranscriptError, match=".txt, .md, .docx"):
        notes_from_file("deck.pptx", b"PK")
    with pytest.raises(TranscriptError, match="Word document could not be read"):
        notes_from_file("bad.docx", b"not a zip")
    with pytest.raises(TranscriptError, match="UTF-8"):
        notes_from_file("x.txt", b"\xff\xfe\xfa")
    with pytest.raises(TranscriptError, match="No speech"):
        vtt_to_notes("WEBVTT\n\nNOTE nothing here\n")
    with pytest.raises(TranscriptError, match="2 MB"):
        notes_from_file("big.txt", b"x" * 2_000_001)
